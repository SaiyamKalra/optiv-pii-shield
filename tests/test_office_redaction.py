# tests/test_office_redaction.py
"""DOCX/PPTX redaction: text masked in every part, PII in images blacked out,
namespaces kept intact, and the rescan verifies the saved file."""

import io
import re
import sys
from pathlib import Path
from zipfile import ZipFile

import docx
import pptx
from lxml import etree
from PIL import Image

sys.path.append(str(Path(__file__).parent.parent))

from models.schemas import (
    BoundingBox,
    DetectedEntityDetail,
    ExtractionInput,
    PIIOutput,
    UnreadRegion,
)
from redaction.office_pipeline import OfficeRedactionPipeline, scan_chunks
from redaction.office_redactor import image_boxes_to_redact
from redaction.office_xml_utils import (
    MASK_CHAR,
    mask_text,
    redact_image,
    redact_xml,
    value_pattern,
)

SECRETS = ["John Smith", "john.smith@cadencefg.com", "EMP-38104"]


class FakePII:
    """Finds the known SECRETS in raw_text (stands in for Presidio + spaCy)."""

    def __init__(self, values=SECRETS):
        self.values = values
        self.calls = 0

    def process_document_chunks(self, chunks):
        self.calls += 1
        outputs = []
        for chunk in chunks:
            found = [
                DetectedEntityDetail(token="[X]", entity_type="TEST", original_value=m.group(),
                                     start=m.start(), end=m.end(), score=1.0,
                                     page_number=chunk.page_number)
                for v in self.values
                for m in re.finditer(re.escape(v), chunk.raw_text)
            ]
            outputs.append(PIIOutput(file_id=chunk.file_id, file_name=chunk.file_name,
                                     page_number=chunk.page_number, status="SUCCESS",
                                     sanitized_text="", detected_entities=found,
                                     verification_passed=True))
        return outputs


def _no_ocr_extractor(path):
    from extraction.normalizer import extract_document
    return extract_document(path)


def _docx(path: Path):
    d = docx.Document()
    p = d.add_paragraph("Contact ")
    p.add_run("John").bold = True  # PII split across runs
    p.add_run(" Smith at john.smith@cadencefg.com, staff id EMP-38104.")
    d.add_paragraph("Annual review by the risk team.")
    table = d.add_table(rows=1, cols=2)
    table.cell(0, 0).text = "Owner"
    table.cell(0, 1).text = "John Smith"
    d.core_properties.author = "John Smith"
    d.save(path)


def _pptx(path: Path):
    prs = pptx.Presentation()
    slide = prs.slides.add_slide(prs.slide_layouts[5])
    slide.shapes.title.text = "Owner: John Smith"
    box = slide.shapes.add_textbox(pptx.util.Inches(1), pptx.util.Inches(2),
                                   pptx.util.Inches(6), pptx.util.Inches(1))
    box.text_frame.text = "Email john.smith@cadencefg.com"
    slide.notes_slide.notes_text_frame.text = "Ask EMP-38104 before sharing"
    prs.save(path)


def _all_text(path: Path) -> str:
    with ZipFile(path) as z:
        return "".join(z.read(n).decode("utf-8", "ignore") for n in z.namelist()
                       if n.endswith((".xml", ".rels")))


# --- XML masking -------------------------------------------------------------------------

def test_value_pattern_matches_whole_words_and_any_whitespace():
    matcher = value_pattern(["Ann Lee"])
    assert matcher.spans("Ann\n  Lee") == [(0, 9)]
    assert matcher.spans("Annual Leeway") == []


def test_overlapping_values_are_both_masked():
    masked, _ = mask_text("Retain the Vendor now", value_pattern(["Retain the", "the Vendor"]))
    assert masked == MASK_CHAR * 17 + " now"


def test_redact_xml_keeps_namespace_prefixes_and_masks_split_runs():
    xml = (b'<?xml version="1.0" encoding="UTF-8" standalone="yes"?>'
           b'<w:document xmlns:w="http://schemas.openxmlformats.org/wordprocessingml/2006/main" '
           b'xmlns:mc="http://schemas.openxmlformats.org/markup-compatibility/2006" '
           b'xmlns:w14="http://schemas.microsoft.com/office/word/2010/wordml" mc:Ignorable="w14">'
           b'<w:body><w:p><w:r><w:t>Hi Jo</w:t></w:r><w:r><w:t>hn Smith!</w:t></w:r></w:p></w:body>'
           b'</w:document>')
    out, count = redact_xml(xml, value_pattern(["John Smith"]))

    assert count == 1
    text = out.decode("utf-8")
    assert 'xmlns:w14="' in text and 'mc:Ignorable="w14"' in text  # ElementTree drops these
    assert "<w:t>" in text and "standalone" in text
    assert "".join(etree.fromstring(out).itertext()) == "Hi " + MASK_CHAR * 10 + "!"


def test_unchanged_part_is_byte_identical():
    xml = b'<?xml version="1.0"?><a xmlns="urn:x"><p><t>nothing here</t></p></a>'
    assert redact_xml(xml, value_pattern(["John Smith"])) == (xml, 0)


# --- images ------------------------------------------------------------------------------

def test_redact_image_blacks_out_box_and_whole_image():
    buf = io.BytesIO()
    Image.new("RGB", (100, 50), "white").save(buf, format="PNG")

    boxed = Image.open(io.BytesIO(redact_image(buf.getvalue(), [(10, 10, 20, 10)])[0]))
    assert boxed.format == "PNG"
    assert boxed.convert("RGB").getpixel((15, 15)) == (0, 0, 0)
    assert boxed.convert("RGB").getpixel((80, 40)) == (255, 255, 255)

    whole = Image.open(io.BytesIO(redact_image(buf.getvalue(), [None])[0]))
    assert whole.convert("RGB").getpixel((99, 49)) == (0, 0, 0)


def test_image_targets_use_pixel_boxes_and_unread_regions():
    text = "id EMP-38104"
    chunk = ExtractionInput(
        file_id="f", file_name="x.pptx", page_number=1, raw_text=text,
        bounding_boxes=[
            BoundingBox(page=1, text="id", x=1, y=1, w=1, h=1, char_start=0, char_end=2,
                        source="ocr", image_ref="/ppt/media/image1.png", image_box=(5, 5, 10, 10)),
            # Slide coordinates, but image_box gives the pixels.
            BoundingBox(page=1, text="EMP-38104", x=300, y=200, w=50, h=10, char_start=3,
                        char_end=12, source="ocr", image_ref="/ppt/media/image1.png",
                        image_box=(20, 5, 40, 10)),
        ],
        unread_regions=[
            UnreadRegion(page=1, x=0, y=0, w=0, h=0, coord_space="image_px",
                         image_ref="/word/media/image9.emf", reason="image_unreadable"),
        ],
    )
    result = FakePII(["EMP-38104"]).process_document_chunks([chunk])

    targets = image_boxes_to_redact([chunk], result)
    assert targets == {"/ppt/media/image1.png": [(20, 5, 40, 10)], "/word/media/image9.emf": [None]}


# --- end to end --------------------------------------------------------------------------

def test_docx_is_redacted_and_verified(tmp_path):
    src, out = tmp_path / "in.docx", tmp_path / "out.docx"
    _docx(src)

    result = OfficeRedactionPipeline(pii_pipeline=FakePII(), extractor=_no_ocr_extractor).process(src, out)

    assert result["status"] == "VERIFIED_NO_PII_FOUND", result
    assert result["verification_passed"] and result["passes"] == 1
    text = _all_text(out)
    for secret in SECRETS:
        assert secret not in text
    assert "Annual review" in text  # not over-masked
    opened = docx.Document(out)  # still a valid Word file
    assert MASK_CHAR in opened.paragraphs[0].text
    assert opened.core_properties.author == MASK_CHAR * 10  # hidden metadata scanned too


def test_pptx_is_redacted_and_verified(tmp_path):
    src, out = tmp_path / "in.pptx", tmp_path / "out.pptx"
    _pptx(src)

    result = OfficeRedactionPipeline(pii_pipeline=FakePII(), extractor=_no_ocr_extractor).process(src, out)

    assert result["status"] == "VERIFIED_NO_PII_FOUND", result
    text = _all_text(out)
    for secret in SECRETS:
        assert secret not in text
    slide = pptx.Presentation(out).slides[0]
    assert slide.shapes.title.text.startswith("Owner: " + MASK_CHAR)


def test_pii_left_after_rescan_blocks(tmp_path):
    src, out = tmp_path / "in.docx", tmp_path / "out.docx"
    _docx(src)

    class Stubborn(FakePII):
        """Always reports a finding, so verification can never pass."""

        def process_document_chunks(self, chunks):
            outputs = super().process_document_chunks(chunks)
            outputs[0].detected_entities.append(DetectedEntityDetail(
                token="[X]", entity_type="TEST", original_value="ghost", start=0, end=1,
                score=1.0, page_number=1))
            return outputs

    result = OfficeRedactionPipeline(pii_pipeline=Stubborn(), extractor=_no_ocr_extractor).process(src, out)

    assert result["status"] == "BLOCKED"
    assert not result["verification_passed"]
    assert result["passes"] == 2
    assert any("Rescan still finds PII" in i["reason"] for i in result["inspection_issues"])


def test_scan_chunks_appends_hidden_text_after_visible_text():
    from models.schemas import HiddenContent
    chunk = ExtractionInput(file_id="f", file_name="x.docx", raw_text="visible",
                            hidden_content=[HiddenContent(kind="metadata", text="John Smith")])
    scanned = scan_chunks([chunk])[0]
    assert scanned.raw_text.startswith("visible") and scanned.raw_text.endswith("John Smith")
    assert chunk.raw_text == "visible"  # original untouched


def test_values_are_split_at_table_cells_and_lines():
    from redaction.office_xml_utils import value_pieces

    assert value_pieces("VP InfoSec | H. Park\nAI") == ["VP InfoSec", "H. Park"]


def test_unwritable_image_format_is_saved_as_png_and_relinked(tmp_path, monkeypatch):
    from redaction import office_xml_utils

    buf = io.BytesIO()
    Image.new("RGB", (60, 30), "white").save(buf, format="BMP")
    src, out = tmp_path / "in.docx", tmp_path / "out.docx"
    with ZipFile(src, "w") as z:
        z.writestr("[Content_Types].xml",
                   '<?xml version="1.0" encoding="UTF-8" standalone="yes"?>'
                   '<Types xmlns="http://schemas.openxmlformats.org/package/2006/content-types">'
                   '<Default Extension="bmp" ContentType="image/bmp"/>'
                   '<Override PartName="/word/media/image1.bmp" ContentType="image/bmp"/></Types>')
        z.writestr("word/_rels/document.xml.rels",
                   '<?xml version="1.0" encoding="UTF-8" standalone="yes"?>'
                   '<Relationships xmlns="http://schemas.openxmlformats.org/package/2006/relationships">'
                   '<Relationship Id="rId1" Type="image" Target="media/image1.bmp"/></Relationships>')
        z.writestr("word/media/image1.bmp", buf.getvalue())
    # Stand in for EMF/WMF: readable by Pillow, but treated as not writable.
    monkeypatch.setattr(office_xml_utils, "SAVABLE_IMAGE_FORMATS", {"PNG"})

    stats = office_xml_utils.redact_office_archive(
        src, out, [], {"/word/media/image1.bmp": [(5, 5, 10, 10)]})

    assert stats["renamed_images"] == {"/word/media/image1.bmp": "/word/media/image1_redacted.png"}
    with ZipFile(out) as z:
        assert "word/media/image1.bmp" not in z.namelist()
        img = Image.open(io.BytesIO(z.read("word/media/image1_redacted.png")))
        assert img.format == "PNG" and img.convert("RGB").getpixel((10, 10)) == (0, 0, 0)
        assert b'Target="media/image1_redacted.png"' in z.read("word/_rels/document.xml.rels")
        types = z.read("[Content_Types].xml")
        assert b'Extension="png"' in types and b"image1.bmp" not in types


def test_ner_guesses_left_after_rescan_need_review_not_block(tmp_path):
    src, out = tmp_path / "in.docx", tmp_path / "out.docx"
    _docx(src)

    class Guessy(FakePII):
        """Like the real NER: tags masked text ("Contact ████") as a name."""

        def process_document_chunks(self, chunks):
            outputs = super().process_document_chunks(chunks)
            for chunk, output in zip(chunks, outputs):
                m = re.search(r"Contact █+", chunk.raw_text)
                if m:
                    output.detected_entities.append(DetectedEntityDetail(
                        token="[P]", entity_type="PERSON", original_value=m.group(),
                        start=m.start(), end=m.end(), score=0.6, page_number=1))
            return outputs

    result = OfficeRedactionPipeline(pii_pipeline=Guessy(), extractor=_no_ocr_extractor).process(src, out)

    assert result["status"] == "REVIEW_REQUIRED", result
    assert result["passes"] == 1  # guesses about masked text are not masked again
    assert result["verification_passed"] and not result["inspection_issues"]
    assert any("guesses" in w["reason"] for w in result["inspection_warnings"])


def test_value_spanning_paragraphs_of_a_table_cell_is_masked():
    # PPTX cell with "T. Lindqvist" and "CIO" as two paragraphs; extraction reads
    # "T. Lindqvist CIO", so that is the value the detector returns.
    xml = (b'<a:tc xmlns:a="http://schemas.openxmlformats.org/drawingml/2006/main"><a:txBody>'
           b'<a:p><a:r><a:t>T. Lindqvist</a:t></a:r></a:p><a:p><a:r><a:t>CIO</a:t></a:r></a:p>'
           b'</a:txBody></a:tc>')
    out, count = redact_xml(xml, value_pattern(["T. Lindqvist CIO"]))

    assert count == 1
    assert [t.text for t in etree.fromstring(out).iter("{*}t")] == [MASK_CHAR * 12, MASK_CHAR * 3]


def test_ner_values_match_case_sensitively_identifiers_do_not():
    from redaction.office_xml_utils import pii_terms

    found = [DetectedEntityDetail(token="[O]", entity_type="ORGANIZATION", original_value="AND",
                                  start=0, end=3, score=0.5),
             DetectedEntityDetail(token="[E]", entity_type="EMAIL_ADDRESS",
                                  original_value="Jo@X.com", start=0, end=8, score=1.0)]
    results = [PIIOutput(file_id="f", file_name="x.docx", status="SUCCESS", sanitized_text="",
                         detected_entities=found, verification_passed=True)]
    masked, _ = mask_text("Contact AND Action and jo@x.com", value_pattern(pii_terms(results)))
    assert masked == f"Contact {MASK_CHAR * 3} Action and {MASK_CHAR * 8}"


def test_line_break_separates_words_for_matching():
    xml = (b'<w:p xmlns:w="http://schemas.openxmlformats.org/wordprocessingml/2006/main">'
           b'<w:r><w:t>CD_Staffing</w:t></w:r><w:r><w:t xml:space="preserve"> Vendors</w:t></w:r>'
           b'<w:r><w:br/><w:t>5. Next</w:t></w:r></w:p>')
    out, count = redact_xml(xml, value_pattern(["CD_Staffing Vendors"]))

    assert count == 1
    assert "".join(etree.fromstring(out).itertext()) == MASK_CHAR * 19 + "5. Next"
