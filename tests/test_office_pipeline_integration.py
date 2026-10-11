import zipfile

import pytest
from docx import Document
from pptx import Presentation
from pptx.util import Inches

from redaction.office_pipeline import OfficeRedactionPipeline

PII_EMAIL = "alice@example.com"


def create_docx(path):
    doc = Document()
    doc.add_heading("Employee Record", level=1)
    doc.add_paragraph(f"Contact: {PII_EMAIL}")

    # Add hidden content through a Word comment-like XML part.
    doc.save(path)

    with zipfile.ZipFile(path, "a") as archive:
        archive.writestr(
            "word/comments.xml",
            f"""<?xml version="1.0" encoding="UTF-8"?>
            <w:comments
            xmlns:w="http://schemas.openxmlformats.org/wordprocessingml/2006/main">
            <w:comment w:id="0" w:author="Reviewer">
                <w:p><w:r><w:t>Confidential contact: {PII_EMAIL}</w:t></w:r></w:p>
            </w:comment>
            </w:comments>""",
        )


def create_pptx(path):
    prs = Presentation()
    slide = prs.slides.add_slide(prs.slide_layouts[6])

    textbox = slide.shapes.add_textbox(Inches(1), Inches(1), Inches(6), Inches(1))
    textbox.text = f"Contact: {PII_EMAIL}"

    # Speaker notes are hidden from the normal slide view.
    slide.notes_slide.notes_text_frame.text = f"Private speaker note: {PII_EMAIL}"
    prs.save(path)


def archive_text(path):
    """Read XML text and attributes from an Office archive."""
    parts = []
    with zipfile.ZipFile(path) as archive:
        for name in archive.namelist():
            if name.endswith((".xml", ".rels")):
                parts.append(archive.read(name).decode("utf-8", errors="ignore"))
    return "\n".join(parts)


@pytest.mark.parametrize(
    ("extension", "creator"),
    [
        (".docx", create_docx),
        (".pptx", create_pptx),
    ],
)
def test_office_pipeline_redacts_visible_and_hidden_pii(tmp_path, extension, creator):
    source = tmp_path / f"source{extension}"
    output = tmp_path / f"sanitized{extension}"
    creator(source)

    pipeline = OfficeRedactionPipeline()
    result = pipeline.process(source, output)

    assert result["status"] == "VERIFIED_NO_PII_FOUND", result
    assert result["verification_passed"] is True
    assert output.exists()
    assert output.stat().st_size > 0

    # Check the actual saved Office package, not only the pipeline result.
    content = archive_text(output).lower()
    assert PII_EMAIL.lower() not in content


@pytest.mark.parametrize(
    ("extension", "creator"),
    [
        (".docx", create_docx),
        (".pptx", create_pptx),
    ],
)
def test_office_pipeline_blocks_documents_with_images(tmp_path, extension, creator):
    from PIL import Image, ImageDraw

    source = tmp_path / f"image_source{extension}"
    output = tmp_path / f"image_sanitized{extension}"
    creator(source)

    # Create an image containing PII. Image redaction is not implemented.
    image_path = tmp_path / "pii_image.png"
    image = Image.new("RGB", (600, 100), "white")
    draw = ImageDraw.Draw(image)
    draw.text((10, 20), PII_EMAIL, fill="black")
    image.save(image_path)

    if extension == ".docx":
        doc = Document(source)
        doc.add_picture(str(image_path))
        doc.save(source)
    else:
        prs = Presentation(source)
        slide = prs.slides[0]
        slide.shapes.add_picture(str(image_path), Inches(1), Inches(2), width=Inches(5))
        prs.save(source)

    result = OfficeRedactionPipeline().process(source, output)

    assert result["status"] == "BLOCKED", result
    assert result["verification_passed"] is False
    assert not output.exists()


def test_office_xml_redaction_does_not_modify_source(tmp_path):
    source = tmp_path / "source.docx"
    output = tmp_path / "sanitized.docx"
    create_docx(source)

    original_bytes = source.read_bytes()
    OfficeRedactionPipeline().process(source, output)

    assert source.read_bytes() == original_bytes
