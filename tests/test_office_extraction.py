# tests/test_office_extraction.py

import io
import sys
from pathlib import Path

import pytest
from docx import Document
from docx.enum.text import WD_BREAK
from docx.shared import Inches
from PIL import Image, ImageDraw, ImageFont
from pptx import Presentation
from pptx.util import Inches as PptxInches
from pptx.util import Pt

sys.path.append(str(Path(__file__).parent.parent))

from extraction.docx_extractor import extract_docx
from extraction.normalizer import extract_document, load_image
from extraction.pptx_extractor import extract_pptx
from extraction.validation import DocumentValidationError
from ocr.ocr_engine import OcrEngine

requires_tesseract = pytest.mark.skipif(
    not OcrEngine().is_available(), reason="Tesseract binary not installed"
)
NO_OCR = OcrEngine(tesseract_cmd="does-not-exist.exe")


def _text_png(text: str, size=(900, 120), mode="RGB") -> bytes:
    bg = (255, 255, 255, 0) if mode == "RGBA" else "white"
    img = Image.new(mode, size, bg)
    ImageDraw.Draw(img).text((20, 30), text, fill="black", font=ImageFont.load_default(size=40))
    buf = io.BytesIO()
    img.save(buf, format="PNG")
    return buf.getvalue()


@pytest.fixture
def sample_docx(tmp_path) -> Path:
    doc = Document()
    doc.sections[0].header.paragraphs[0].text = "Confidential header"
    doc.add_paragraph("Contact Jane Doe at jane@example.com")
    table = doc.add_table(rows=2, cols=2)
    table.cell(0, 0).text = "Name"
    table.cell(0, 1).text = "Staff ID"
    table.cell(1, 0).text = "Rafael Duarte"
    table.cell(1, 1).text = "EMP-40944"
    doc.add_paragraph().add_run().add_break(WD_BREAK.PAGE)
    doc.add_paragraph("Second page text")
    doc.add_picture(io.BytesIO(_text_png("Email r.duarte@example.com")), width=Inches(5))
    doc.add_picture(io.BytesIO(_text_png("x", size=(40, 40))), width=Inches(0.2))
    path = tmp_path / "sample.docx"
    doc.save(path)
    return path


@pytest.fixture
def sample_pptx(tmp_path) -> Path:
    prs = Presentation()
    blank = prs.slide_layouts[6]

    slide = prs.slides.add_slide(blank)
    box = slide.shapes.add_textbox(PptxInches(1), PptxInches(1), PptxInches(6), PptxInches(1))
    box.text_frame.text = "Escalate to Marguerite Vossberg"
    table = slide.shapes.add_table(2, 2, PptxInches(1), PptxInches(3), PptxInches(6), PptxInches(1)).table
    table.cell(0, 0).text = "Mobile"
    table.cell(0, 1).text = "+1 (555) 0114"
    table.cell(1, 0).text = "Email"
    table.cell(1, 1).text = "m.vossberg@cadencefg.example"
    slide.notes_slide.notes_text_frame.text = "Speaker note: call Tobias Lindqvist"

    slide = prs.slides.add_slide(blank)
    slide.shapes.add_picture(
        io.BytesIO(_text_png("Passport TK4820913")), PptxInches(1), PptxInches(1), width=PptxInches(6)
    )
    slide.shapes.add_picture(
        io.BytesIO(_text_png("x", size=(40, 40))), PptxInches(8), PptxInches(0.2), width=Pt(20)
    )

    prs.slides.add_slide(blank)  # empty slide
    path = tmp_path / "sample.pptx"
    prs.save(path)
    return path


# --- shared behaviour ----------------------------------------------------------------

@pytest.mark.parametrize("fixture", ["sample_docx", "sample_pptx"])
def test_metadata_identical_and_pages_in_order(fixture, request):
    path = request.getfixturevalue(fixture)
    chunks = extract_document(path, source_url="storage/x", ocr_engine=NO_OCR)

    assert [c.page_number for c in chunks] == list(range(1, len(chunks) + 1))
    assert {c.file_id for c in chunks} == {chunks[0].file_id}
    assert {c.file_name for c in chunks} == {path.name}
    assert {c.source_url for c in chunks} == {"storage/x"}


def test_unsupported_file_type(tmp_path):
    path = tmp_path / "notes.txt"
    path.write_text("hello")
    with pytest.raises(DocumentValidationError) as err:
        extract_document(path)
    assert err.value.reason == "unsupported_type"


def test_load_image_rejects_non_images():
    assert load_image(b"<svg xmlns='http://www.w3.org/2000/svg'></svg>") is None


def test_load_image_flattens_transparency_onto_white():
    img = load_image(_text_png("x", mode="RGBA"))

    assert img.mode == "RGB"
    assert img.getpixel((0, 0)) == (255, 255, 255)


# --- DOCX ---------------------------------------------------------------------------

def test_docx_text_tables_and_header(sample_docx):
    chunks = extract_docx(sample_docx, ocr_engine=NO_OCR)
    page1 = chunks[0].raw_text

    assert "Confidential header" in page1
    assert "jane@example.com" in page1
    assert "Rafael Duarte | EMP-40944" in page1  # a table row stays on one line


def test_docx_splits_on_page_break(sample_docx):
    chunks = extract_docx(sample_docx, ocr_engine=NO_OCR)

    assert len(chunks) == 2
    assert "Second page text" in chunks[1].raw_text
    assert "Second page text" not in chunks[0].raw_text


def test_docx_images_reported_when_ocr_unavailable(sample_docx):
    page2 = extract_docx(sample_docx, ocr_engine=NO_OCR)[1]

    assert page2.ocr_used is False
    assert page2.images_skipped_count == 2  # the text image (no OCR) + the tiny icon


@requires_tesseract
def test_docx_image_text_is_ocrd_with_image_coordinates(sample_docx):
    page2 = extract_docx(sample_docx)[1]

    assert page2.extraction_method == "hybrid"
    assert page2.images_ocr_count == 1
    assert page2.images_skipped_count == 1  # tiny icon
    assert "r.duarte@example.com" in page2.raw_text
    ocr_boxes = [b for b in page2.bounding_boxes if b.source == "ocr"]
    assert ocr_boxes
    for b in ocr_boxes:
        assert b.coord_space == "image_px"
        assert b.image_ref.startswith("/word/media/")
        assert page2.raw_text[b.char_start:b.char_end] == b.text


# --- PPTX ---------------------------------------------------------------------------

def test_pptx_text_table_and_notes(sample_pptx):
    slide1 = extract_pptx(sample_pptx, ocr_engine=NO_OCR)[0]

    assert "Marguerite Vossberg" in slide1.raw_text
    assert "Mobile | +1 (555) 0114" in slide1.raw_text
    assert "m.vossberg@cadencefg.example" in slide1.raw_text
    assert "Tobias Lindqvist" in slide1.raw_text  # speaker notes
    assert slide1.extraction_method == "native"


def test_pptx_one_chunk_per_slide_including_empty(sample_pptx):
    chunks = extract_pptx(sample_pptx, ocr_engine=NO_OCR)

    assert len(chunks) == 3
    assert chunks[2].raw_text == ""


@requires_tesseract
def test_pptx_picture_ocr_in_slide_points(sample_pptx):
    slide2 = extract_pptx(sample_pptx)[1]
    picture = (72, 72, 72 + 6 * 72)  # left, top, right in points (1 inch = 72 pt)

    assert slide2.extraction_method == "ocr"
    assert slide2.images_ocr_count == 1
    assert slide2.images_skipped_count == 1  # tiny icon
    assert "TK4820913" in slide2.raw_text
    for b in slide2.bounding_boxes:
        assert b.coord_space == "page_pt"
        assert b.image_ref.startswith("/ppt/media/")
        assert picture[0] <= b.x and b.x + b.w <= picture[2]
        assert b.y >= picture[1]
        assert slide2.raw_text[b.char_start:b.char_end] == b.text


@requires_tesseract
def test_office_output_feeds_pii_pipeline(sample_docx, sample_pptx):
    spacy = pytest.importorskip("spacy")
    if not spacy.util.is_package("en_core_web_lg"):
        pytest.skip("spaCy model en_core_web_lg not installed")
    from pii.pipeline import PIIPipeline

    pipeline = PIIPipeline()
    for path in (sample_docx, sample_pptx):
        results = pipeline.process_document_chunks(extract_document(path))
        assert results
        for r in results:
            assert "@example.com" not in r.sanitized_text
            assert "@cadencefg.example" not in r.sanitized_text
