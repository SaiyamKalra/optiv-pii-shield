import io
import shutil
import sys
from pathlib import Path

import pymupdf
import pytest
from docx import Document
from docx.shared import Inches
from PIL import Image, ImageDraw, ImageFont
from pptx import Presentation
from pptx.util import Inches as PptxInches

sys.path.append(str(Path(__file__).parent.parent))

from extraction import normalizer
from extraction.docx_extractor import extract_docx
from extraction.normalizer import PageTextBuilder, extract_document, image_hash
from extraction.pdf_extractor import extract_pdf
from extraction.pptx_extractor import extract_pptx
from extraction.validation import DocumentValidationError, validate_document
from ocr.ocr_engine import OcrEngine

requires_tesseract = pytest.mark.skipif(
    not OcrEngine().is_available(), reason="Tesseract binary not installed"
)
NO_OCR = OcrEngine(tesseract_cmd="does-not-exist.exe")


def _text_png(text: str, size=(900, 140)) -> bytes:
    img = Image.new("RGB", size, "white")
    ImageDraw.Draw(img).text((20, 40), text, fill="black", font=ImageFont.load_default(size=44))
    buf = io.BytesIO()
    img.save(buf, format="PNG")
    return buf.getvalue()


def _barcode_png(size=(900, 300)) -> bytes:
    img = Image.new("RGB", size, "white")
    draw = ImageDraw.Draw(img)
    draw.text((20, 20), "Badge holder details", fill="black", font=ImageFont.load_default(size=40))
    for i in range(60):
        x = 60 + i * 12
        draw.rectangle([x, 150, x + (3 if i % 3 else 7), 260], fill="black")
    buf = io.BytesIO()
    img.save(buf, format="PNG")
    return buf.getvalue()


def _pdf_with_image(tmp_path, png: bytes, name="img.pdf", rotation=0) -> Path:
    doc = pymupdf.open()
    page = doc.new_page(width=612, height=792)
    page.insert_image(pymupdf.Rect(72, 200, 522, 270 if png != _barcode_png() else 350), stream=png)
    page.set_rotation(rotation)
    path = tmp_path / name
    doc.save(path)
    doc.close()
    return path


def _simple_pdf(path: Path, **save_kwargs) -> Path:
    doc = pymupdf.open()
    doc.new_page().insert_text((72, 72), "hello")
    doc.save(path, **save_kwargs)
    doc.close()
    return path


@pytest.mark.parametrize(
    ("make", "reason"),
    [
        (lambda p: p.write_bytes(b""), "empty"),
        (lambda p: p.write_bytes(_text_png("x")), "unsupported_type"),
        (lambda p: p.write_bytes(b"%PDF-1.7\n" + b"garbage" * 50), "corrupt"),
        (lambda p: p.write_bytes(b"\xd0\xcf\x11\xe0\xa1\xb1\x1a\xe1" + b"\x00" * 600), "encrypted_or_legacy"),
    ],
    ids=["empty", "renamed_png", "truncated_pdf", "ole_container"],
)
def test_validation_rejects_bad_files(tmp_path, make, reason):
    path = tmp_path / "upload.pdf"
    make(path)
    with pytest.raises(DocumentValidationError) as err:
        extract_document(path)
    assert err.value.reason == reason


def test_validation_missing_file(tmp_path):
    with pytest.raises(DocumentValidationError) as err:
        validate_document(tmp_path / "nope.pdf")
    assert err.value.reason == "not_found"


def test_validation_password_protected_pdf(tmp_path):
    path = _simple_pdf(
        tmp_path / "locked.pdf",
        encryption=pymupdf.PDF_ENCRYPT_AES_256, user_pw="secret", owner_pw="owner",
    )
    with pytest.raises(DocumentValidationError) as err:
        extract_pdf(path)
    assert err.value.reason == "encrypted"


def test_validation_zip_that_is_not_office(tmp_path):
    base = tmp_path / "archive"
    (tmp_path / "data").mkdir()
    (tmp_path / "data" / "a.txt").write_text("x")
    shutil.make_archive(str(base), "zip", tmp_path / "data")
    path = base.with_suffix(".docx")
    base.with_suffix(".zip").rename(path)
    with pytest.raises(DocumentValidationError) as err:
        extract_document(path)
    assert err.value.reason == "unsupported_type"


def test_type_detected_by_content_not_extension(tmp_path):
    doc = Document()
    doc.add_paragraph("Contact jane@example.com")
    path = tmp_path / "really_a_docx.pdf"
    doc.save(path)

    chunks = extract_document(path, ocr_engine=NO_OCR)
    assert "jane@example.com" in chunks[0].raw_text


@requires_tesseract
@pytest.mark.parametrize("rotation", [90, 180, 270])
def test_rotated_page_images_are_ocrd_in_unrotated_coordinates(tmp_path, rotation):
    path = _pdf_with_image(tmp_path, _text_png("Passport TK4820913"), rotation=rotation)
    page = extract_pdf(path)[0]

    assert "TK4820913" in page.raw_text
    image_rect = pymupdf.Rect(72, 200, 522, 270)
    for box in page.bounding_boxes:
        assert image_rect.contains(pymupdf.Rect(box.x, box.y, box.x + box.w, box.y + box.h))


def test_pdf_hidden_content(tmp_path):
    doc = pymupdf.open()
    page = doc.new_page()
    page.insert_text((72, 72), "Visible body text")
    page.add_text_annot((100, 100), "Reviewer note: call Jane on 555-0101").set_info(title="Jane Roe")
    widget = pymupdf.Widget()
    widget.field_name = "employee_name"
    widget.field_type = pymupdf.PDF_WIDGET_TYPE_TEXT
    widget.field_value = "Rafael Duarte"
    widget.rect = pymupdf.Rect(72, 300, 300, 320)
    page.add_widget(widget)
    page.insert_link({"kind": pymupdf.LINK_URI, "from": pymupdf.Rect(72, 400, 200, 420),
                      "uri": "mailto:r.duarte@example.com"})
    doc.set_metadata({"author": "Miriam Achebe", "title": "Risk policy"})
    doc.set_toc([[1, "Contacts for Naomi Osadebe", 1]])
    doc.embfile_add("roster.csv", b"name,ssn\nx,900-00-0000", filename="roster.csv", desc="staff roster")
    path = tmp_path / "hidden.pdf"
    doc.save(path)
    doc.close()

    page1 = extract_pdf(path, ocr_engine=NO_OCR)[0]
    kinds = {(h.kind, h.text) for h in page1.hidden_content}
    texts = " ".join(h.text for h in page1.hidden_content)

    assert ("metadata", "Miriam Achebe") in kinds
    assert "Reviewer note: call Jane on 555-0101" in texts and "Jane Roe" in texts
    assert ("form_field", "Rafael Duarte") in kinds
    assert ("link", "mailto:r.duarte@example.com") in kinds
    assert ("bookmark", "Contacts for Naomi Osadebe") in kinds
    assert any(h.kind == "embedded_file" and "roster.csv" in h.text for h in page1.hidden_content)
    assert "Reviewer note" not in page1.raw_text


def test_docx_hidden_content(tmp_path):
    doc = Document()
    doc.core_properties.author = "Ifeoma Adeyemi"
    doc.core_properties.comments = "Draft prepared for Hiroshi Yamagata"
    paragraph = doc.add_paragraph("Escalate to the CRO.")
    doc.add_comment(paragraph.runs, text="Use her mobile +1 917 555 0146", author="Sofia Lang")
    deleted = paragraph._p.makeelement(
        "{http://schemas.openxmlformats.org/wordprocessingml/2006/main}del", {}
    )
    run = deleted.makeelement("{http://schemas.openxmlformats.org/wordprocessingml/2006/main}r", {})
    text = run.makeelement("{http://schemas.openxmlformats.org/wordprocessingml/2006/main}delText", {})
    text.text = "SSN 900-55-2043"
    run.append(text)
    deleted.append(run)
    paragraph._p.append(deleted)
    doc.add_picture(io.BytesIO(_text_png("x")), width=Inches(3))
    doc.inline_shapes[0]._inline.docPr.set("descr", "Photo of Naomi Osadebe")
    path = tmp_path / "hidden.docx"
    doc.save(path)

    page1 = extract_docx(path, ocr_engine=NO_OCR)[0]
    kinds = {(h.kind, h.text) for h in page1.hidden_content}

    assert ("metadata", "Ifeoma Adeyemi") in kinds
    assert ("metadata", "Draft prepared for Hiroshi Yamagata") in kinds
    assert ("comment", "Sofia Lang: Use her mobile +1 917 555 0146") in kinds
    assert ("deleted_text", "SSN 900-55-2043") in kinds
    assert ("alt_text", "Photo of Naomi Osadebe") in kinds
    assert "900-55-2043" not in page1.raw_text


def test_pptx_hidden_content(tmp_path):
    prs = Presentation()
    prs.core_properties.author = "Adaeze Okonkwo-Bell"
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    pic = slide.shapes.add_picture(io.BytesIO(_text_png("x")), PptxInches(1), PptxInches(1))
    pic._element.nvPicPr.cNvPr.set("descr", "Headshot of Solveig Aarnio")
    hidden = prs.slides.add_slide(prs.slide_layouts[6])
    hidden.shapes.add_textbox(PptxInches(1), PptxInches(1), PptxInches(4), PptxInches(1)).text_frame.text = "Draft"
    hidden._element.set("show", "0")
    path = tmp_path / "hidden.pptx"
    prs.save(path)

    slides = extract_pptx(path, ocr_engine=NO_OCR)
    kinds1 = {(h.kind, h.text) for h in slides[0].hidden_content}

    assert ("metadata", "Adaeze Okonkwo-Bell") in kinds1
    assert ("alt_text", "Headshot of Solveig Aarnio") in kinds1
    assert any(h.kind == "hidden_slide" for h in slides[1].hidden_content)
    assert "Draft" in slides[1].raw_text


@requires_tesseract
def test_unreadable_ink_is_reported_as_unread_region(tmp_path):
    doc = pymupdf.open()
    page = doc.new_page(width=612, height=792)
    image_rect = pymupdf.Rect(72, 200, 522, 350)
    page.insert_image(image_rect, stream=_barcode_png())
    path = tmp_path / "barcode.pdf"
    doc.save(path)
    doc.close()

    page1 = extract_pdf(path)[0]

    assert "Badge holder details" in page1.raw_text
    assert page1.unread_regions, "the bars should be reported as unread"
    for region in page1.unread_regions:
        assert image_rect.contains(pymupdf.Rect(region.x, region.y, region.x + region.w, region.y + region.h))
        assert region.y > 200 + 150 * 0.4


@requires_tesseract
def test_clean_text_image_has_no_unread_regions(tmp_path):
    path = _pdf_with_image(tmp_path, _text_png("Passport TK4820913"))
    assert extract_pdf(path)[0].unread_regions == []


def test_punctuation_does_not_count_as_low_confidence():
    builder = PageTextBuilder(1, low_conf_threshold=60)
    builder.add_word("|", 0, 0, 1, 1, source="ocr", confidence=5)
    builder.add_word("—", 0, 0, 1, 1, source="ocr", confidence=10)
    builder.add_word("EMP-38104", 0, 0, 1, 1, source="ocr", confidence=95)
    data = builder.page_data()

    assert data["low_confidence_word_count"] == 0
    assert data["ocr_mean_confidence"] == 95
    assert "|" in data["raw_text"]


def test_decorative_images_are_not_ocrd(tmp_path, monkeypatch):
    logo = _text_png("ACME LOGO")
    monkeypatch.setattr(normalizer, "decorative_hashes", lambda: frozenset({image_hash(logo)}))
    prs = Presentation()
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    slide.shapes.add_picture(io.BytesIO(logo), PptxInches(1), PptxInches(1), width=PptxInches(4))
    path = tmp_path / "logo.pptx"
    prs.save(path)

    slide1 = extract_pptx(path)[0]

    assert slide1.images_decorative_count == 1
    assert slide1.images_ocr_count == 0
    assert slide1.images_skipped_count == 0
    assert "ACME" not in slide1.raw_text


@requires_tesseract
def test_ocr_failure_is_reported_with_unread_region(tmp_path, monkeypatch):
    def boom(self, img, px_per_inch=None):
        raise RuntimeError("Tesseract process timeout")

    monkeypatch.setattr(OcrEngine, "ocr_image_detailed", boom)
    path = _pdf_with_image(tmp_path, _text_png("Passport TK4820913"))
    page1 = extract_pdf(path)[0]

    assert page1.images_skipped_count == 1
    assert any("timeout" in e for e in page1.extraction_errors)
    assert [r.reason for r in page1.unread_regions] == ["image_error"]
